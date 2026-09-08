import openpyxl
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, Http404
from django.db import transaction
from django.core.mail import EmailMessage
from django.conf import settings
from django.core.paginator import Paginator
from django.db.models import Q
import os

from core.decorators import admin_required
from core.utils import generate_college_qr, send_college_qr_email
from core.models import AuditLog
from .models import College
from .forms import CollegeForm

@login_required
@admin_required
def college_list(request):
    """
    Renders the list of colleges with search, filters and pagination.
    """
    queryset = College.objects.all()
    
    # Extract Search Parameters
    search_query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('status', '').strip()
    
    if search_query:
        queryset = queryset.filter(
            Q(college_name__icontains=search_query) |
            Q(college_code__icontains=search_query) |
            Q(email__icontains=search_query)
        )
        
    if status_filter:
        queryset = queryset.filter(status=status_filter)
        
    paginator = Paginator(queryset, 10)  # Show 10 colleges per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    context = {
        'page_obj': page_obj,
        'q': search_query,
        'status': status_filter,
    }
    return render(request, 'colleges/list.html', context)

@login_required
@admin_required
def college_detail(request, pk):
    """
    Renders details of a specific college and its scan logs history.
    """
    college = get_object_or_404(College, pk=pk)
    scan_logs = college.scan_logs.all()[:20]  # Show recent 20 logs
    context = {
        'college': college,
        'scan_logs': scan_logs,
    }
    return render(request, 'colleges/detail.html', context)

@login_required
@admin_required
def college_create(request):
    """
    Handles creating a new college record and generating its secure QR code.
    """
    if request.method == 'POST':
        form = CollegeForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                college = form.save(commit=False)
                college.used_pass = 0
                college.save()
                
                # Generate QR code immediately after save
                generate_college_qr(college)
            
            # Log to audit trail
            AuditLog.objects.create(
                user=request.user,
                action="Create College",
                details=f"Created college '{college.college_name}' ({college.college_code}) with generated QR code.",
                ip_address=request.META.get('REMOTE_ADDR')
            )
            messages.success(request, f"College '{college.college_name}' created successfully with QR code. You can send the E-QR pass email manually.")
            return redirect('college_detail', pk=college.pk)
    else:
        form = CollegeForm()
    return render(request, 'colleges/form.html', {'form': form, 'title': 'Create College'})

@login_required
@admin_required
def college_edit(request, pk):
    """
    Handles editing an existing college record.
    """
    college = get_object_or_404(College, pk=pk)
    if request.method == 'POST':
        form = CollegeForm(request.POST, instance=college)
        if form.is_valid():
            # If college code changes, we regenerate the QR code image filename accordingly
            old_code = college.college_code
            college = form.save()
            if old_code != college.college_code:
                generate_college_qr(college)
            messages.success(request, f"College '{college.college_name}' details updated successfully.")
            return redirect('college_detail', pk=college.pk)
    else:
        form = CollegeForm(instance=college)
    return render(request, 'colleges/form.html', {'form': form, 'title': 'Edit College', 'college': college})

@login_required
@admin_required
def college_delete(request, pk):
    """
    Handles deleting a college record.
    """
    college = get_object_or_404(College, pk=pk)
    college_name = college.college_name
    
    # Clean up file from media folder
    if college.qr_image and os.path.exists(college.qr_image.path):
        try:
            os.remove(college.qr_image.path)
        except Exception as e:
            pass
            
    college.delete()
    messages.success(request, f"College '{college_name}' deleted successfully.")
    return redirect('college_list')

@login_required
@admin_required
def college_qr_download(request, pk):
    """
    Triggers attachment file download for the college's generated QR code image.
    """
    college = get_object_or_404(College, pk=pk)
    if not college.qr_image:
        generate_college_qr(college)
        
    file_path = college.qr_image.path
    if os.path.exists(file_path):
        with open(file_path, 'rb') as fh:
            response = HttpResponse(fh.read(), content_type="image/png")
            response['Content-Disposition'] = f'attachment; filename="QR_{college.college_code}.png"'
            return response
    raise Http404("QR Code image not found.")

@login_required
@admin_required
def college_qr_email(request, pk):
    """
    Sends the college QR Code to the contact person via email using professional HTML template.
    """
    college = get_object_or_404(College, pk=pk)
    try:
        send_college_qr_email(college, request=request, action="Email QR Pass")
        messages.success(request, f"E-QR Pass successfully sent to {college.email}.")
    except Exception as e:
        messages.error(request, f"Failed to send email: {str(e)}")
        
    return redirect('college_detail', pk=college.pk)

@login_required
@admin_required
def colleges_import_excel(request):
    """
    Imports college records from an uploaded Excel file.
    Creates college records and triggers QR code generation.
    """
    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        
        try:
            wb = openpyxl.load_workbook(excel_file)
            sheet = wb.active
            
            imported_count = 0
            skipped_count = 0
            
            # Start loop from row 2 (row 1 is header: Name, Code, Contact, Email, Phone, MaxPass)
            for row in sheet.iter_rows(min_row=2, values_only=True):
                # Skip completely empty rows
                if not row or not any(row):
                    continue
                
                college_name = row[0]
                college_code = row[1]
                email = row[3]
                max_pass = int(row[5]) if len(row) > 5 and row[5] is not None else 2
                
                if not college_name or not college_code or not email:
                    skipped_count += 1
                    continue
                
                # Check for duplicates
                if College.objects.filter(college_code=str(college_code).strip()).exists():
                    skipped_count += 1
                    continue
                    
                with transaction.atomic():
                    college = College.objects.create(
                        college_name=str(college_name).strip(),
                        college_code=str(college_code).strip(),
                        email=str(email).strip(),
                        max_pass=max_pass,
                        used_pass=0,
                        status=College.Status.ACTIVE
                    )
                    # Generate QR Code image
                    generate_college_qr(college)
                    imported_count += 1
            
            messages.success(request, f"Excel Import Completed! {imported_count} colleges imported successfully. {skipped_count} skipped (duplicates or empty fields).")
        except Exception as e:
            messages.error(request, f"Error processing Excel file: {str(e)}")
            
    return redirect('college_list')

@login_required
@admin_required
def colleges_export_excel(request):
    """
    Generates and downloads an Excel sheet of all registered colleges and their entry status.
    """
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="colleges_report.xlsx"'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Colleges Entry Summary"
    
    # Headers
    headers = [
        "College Name", "College Code", "Email", "Max Passes", "Used Passes", "Remaining Passes", "Status"
    ]
    ws.append(headers)
    
    for college in College.objects.all():
        row = [
            college.college_name,
            college.college_code,
            college.email,
            college.max_pass,
            college.used_pass,
            college.remaining_passes(),
            college.status
        ]
        ws.append(row)
        
    wb.save(response)
    return response

@login_required
@admin_required
def colleges_bulk_email_qr(request):
    """
    Sends the generated QR Passes to all active colleges via email using professional HTML template.
    """
    active_colleges = College.objects.filter(status=College.Status.ACTIVE)
    total = active_colleges.count()
    
    if total == 0:
        messages.warning(request, "No active colleges found to email passes.")
        return redirect('college_list')
        
    sent_count = 0
    failed_count = 0
    
    for college in active_colleges:
        try:
            send_college_qr_email(college, request=request, action="Bulk Email QR Pass")
            sent_count += 1
        except Exception as e:
            failed_count += 1
            
    if sent_count > 0:
        messages.success(request, f"E-QR passes successfully emailed to {sent_count} colleges.")
    if failed_count > 0:
        messages.error(request, f"Failed to email passes to {failed_count} colleges.")
        
    return redirect('college_list')

